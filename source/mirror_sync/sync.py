"""Server-only, read-only Drive XLSX -> validated immutable RTDB mirror.

Run as a dedicated Cloud Run Job service account; Cloud Scheduler invokes it
every minute. Authentication uses Application Default Credentials with scopes;
never store service account keys in source or client configuration.
"""
from __future__ import annotations
import hashlib
import io
import json
import os
import re
import sys
import unicodedata
from datetime import datetime, timezone
from urllib.parse import quote

import google.auth
from google.auth.transport.requests import AuthorizedSession
from openpyxl import load_workbook

SCOPES = (
    "https://www.googleapis.com/auth/drive.readonly",
    "https://www.googleapis.com/auth/firebase.database",
    "https://www.googleapis.com/auth/userinfo.email",
)
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
FIELDS = ("content", "department", "type", "partyLevel", "progressParty", "progressBuild",
          "progressCouncil", "specialist", "room", "vpLeader", "ubndLeader",
          "specificProgress", "note")


def simple(value):
    value = str(value or "").strip().replace("đ", "d").replace("Đ", "D")
    return "".join(c for c in unicodedata.normalize("NFD", value)
                   if unicodedata.category(c) != "Mn").casefold()


def cell_text(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y")
    return str(value).strip()


def parse_workbook(data: bytes, sheet_name: str, source: str):
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise ValueError("Không tìm thấy sheet được chỉ định chính xác: " + sheet_name)
    sheet = workbook[sheet_name]
    raw = list(sheet.iter_rows(values_only=True))
    if not raw or len(raw) > 8000:
        raise ValueError("Workbook quá lớn hoặc không có hàng.")
    header_index = next((i for i, row in enumerate(raw[:10])
                         if len(row) >= 15 and "noi dung" in simple(row[2])
                         and "so" in simple(row[3])), -1)
    if header_index < 0:
        raise ValueError("Sai nguồn: không tìm thấy header cột C/D.")
    header = raw[header_index]
    need = {4: "loai vb", 5: "cap", 6: "tien do xin y kien", 7: "trinh hdnd",
            8: "trinh hdnd", 13: "tien do hien tai", 14: "ghi chu"}
    for col, term in need.items():
        if len(header) <= col or term not in simple(header[col]):
            raise ValueError(f"Sai schema Excel A–O ở cột {col+1}: {term}")
    result = []
    warnings = []
    for pos, row in enumerate(raw[header_index + 1:], start=header_index + 2):
        values = [cell_text(row[i] if i < len(row) else None) for i in range(15)]
        code, ordinal, content = values[:3]
        if not any(values):
            continue
        if not code and not ordinal and not content:
            continue
        if not code and re.fullmatch(r"[IVXLCDM]+\.?", ordinal, re.I) and content:
            continue
        if not content:
            warnings.append(f"Row {pos}: blank content")
        stt = pos - header_index - 1
        try:
            stt = int(float(ordinal))
        except (ValueError, TypeError):
            pass
        if stt < 1:
            stt = pos - header_index - 1
        t = simple(values[4])
        bc_type = t in ("bc", "bao cao") or "bao cao" in t
        nq_type = t in ("hc", "hanh chinh", "qppl", "nq")
        code_bc, code_nq = code.upper().startswith("BC"), code.upper().startswith(("ID", "NQ"))
        kind = "UNCLASSIFIED" if (bc_type and code_nq) or (nq_type and code_bc) else \
               "BC" if bc_type or code_bc else "NQ" if nq_type or code_nq else "UNCLASSIFIED"
        result.append({
            "recordId": f"{source}:row:{pos}", "displayCode": code or f"Dòng {pos} (thiếu mã)",
            "stt": stt, "documentKind": kind,
            **dict(zip(FIELDS, [values[2], values[3], values[4], values[5], values[6],
                                values[7], values[8], values[9], values[10], values[11],
                                values[12], values[13], values[14]])),
            "isDropped": False, "archivedAt": None, "revision": 0,
        })
    if len(result) == 0 or len(result) > 5000:
        raise ValueError("Sai số lượng hàng, không ghi đè bản sao hợp lệ.")
    if len({v["recordId"] for v in result}) != len(result):
        raise ValueError("Trùng recordId.")
    return result, warnings


def required(name: str) -> str:
    value = os.getenv(name, "").strip()
    if not value:
        raise ValueError(f"Missing env {name}")
    return value


def check_response(response, label):
    if not response.ok:
        raise RuntimeError(f"{label}: HTTP {response.status_code}, {response.text[:240]}")
    return response


def run():
    source_id = required("SOURCE_FILE_ID")
    gid = required("SOURCE_GID")
    sheet_name = required("SOURCE_SHEET_NAME")
    base_url = required("RTDB_URL").rstrip("/")
    if not re.fullmatch(r"[-\w]{10,200}", source_id) or not re.fullmatch(r"\d{1,20}", gid):
        raise ValueError("Invalid allowlisted source file ID/gid.")
    if not base_url.startswith("https://") or not base_url.endswith(".firebasedatabase.app"):
        raise ValueError("Refusing untrusted Realtime Database host.")
    credentials, _ = google.auth.default(scopes=SCOPES)
    session = AuthorizedSession(credentials)
    drive_url = f"https://www.googleapis.com/drive/v3/files/{quote(source_id, safe='')}"
    info = check_response(session.get(drive_url, params={"fields":"id,name,mimeType,modifiedTime,size"},
                                      timeout=25), "Drive metadata").json()
    if info.get("id") != source_id or info.get("mimeType") != XLSX_MIME:
        raise ValueError("Expected original Office XLSX; refusing conversion or different file.")
    if int(info.get("size", 0)) > 15_000_000:
        raise ValueError("Workbook exceeds safe 15 MB bound.")
    raw = check_response(session.get(drive_url, params={"alt":"media"}, timeout=45),
                         "Download workbook").content
    digest = hashlib.sha256(raw).hexdigest()
    source_key = source_id + ":" + gid
    rows, warnings = parse_workbook(raw, sheet_name, source_key)
    if len(warnings) > max(5, len(rows) // 10):
        raise ValueError(f"Too many source warnings ({len(warnings)}), refusing mirror update.")
    endpoint = (base_url + "/hdnd/mirrors/" + quote(source_id, safe="") + "/" +
                quote(gid, safe="") + ".json")
    previous_response = check_response(session.get(endpoint, headers={"X-Firebase-ETag":"true"},
                                                   timeout=20), "Read mirror baseline")
    etag = previous_response.headers.get("ETag")
    if not etag:
        raise RuntimeError("Firebase did not return ETag; cannot safely publish.")
    previous = previous_response.json()
    if previous:
        if previous.get("source") != source_key:
            raise ValueError("Mirror belongs to another source.")
        old_at = previous.get("sourceModifiedAt", "")
        if old_at and old_at > info["modifiedTime"]:
            raise ValueError("Source older than latest mirror. Refusing rollback.")
        old_count = len(previous.get("rows") or [])
        if old_count and len(rows) < old_count * 0.8 and \
           os.getenv("ALLOW_MAJOR_ROW_REMOVAL") != "true":
            raise ValueError(f"Large row count drop: {old_count} -> {len(rows)}; manual approval needed.")
        if previous.get("sha256") == digest:
            print(json.dumps({"unchanged":True,"rows":len(rows),"modified":info["modifiedTime"]}))
            return
    mirror = {"schemaVersion":1, "source":source_key, "syncedAt":datetime.now(timezone.utc).isoformat(),
              "sourceModifiedAt":info["modifiedTime"],"sha256":digest,"rows":rows}
    response = session.put(endpoint, headers={"If-Match":etag, "Content-Type":"application/json"},
                           data=json.dumps(mirror,ensure_ascii=False,separators=(",",":")).encode("utf-8"),
                           timeout=35)
    if response.status_code == 412:
        raise RuntimeError("Concurrent sync updated the mirror first; will retry next minute.")
    check_response(response, "Conditional mirror write")
    verified = check_response(session.get(endpoint, timeout=20), "Mirror readback").json()
    if (not verified or verified.get("sha256") != digest or
            len(verified.get("rows") or []) != len(rows)):
        raise RuntimeError("Mirror readback differs; alert operator.")
    print(json.dumps({"published":True,"rows":len(rows),"sha256":digest,
                      "sourceModifiedAt":info["modifiedTime"],"syncedAt":mirror["syncedAt"]}))


if __name__ == "__main__":
    try:
        run()
    except Exception as exc:
        print("MIRROR_SYNC_FAILED: "+str(exc), file=sys.stderr)
        sys.exit(1)
